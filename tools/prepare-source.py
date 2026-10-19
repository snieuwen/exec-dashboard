"""Prepare generic EIX matrices without embedding customer rules or changing input."""
import argparse, datetime, json
from pathlib import Path
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[1]

def prepare(file):
    schema=json.loads((ROOT/'schemas/eix-columns.json').read_text(encoding='utf8'))
    values=load_workbook(file,read_only=True,data_only=True)
    formulas=load_workbook(file,read_only=True,data_only=False)
    aliases={'sw-details':['sw-details','Primary Data'],'hw-details':['hw-details','HW Data','Hardware Data']}
    result={}
    try:
        for name,expected in schema.items():
            candidates=[key for key in aliases[name] if key in values.sheetnames]
            if len(candidates)!=1:raise ValueError(f'Expected one source sheet for {name}; found {candidates}')
            sheet=values[candidates[0]];formula_sheet=formulas[candidates[0]];header=None;mapping=None
            for number,row in enumerate(sheet.iter_rows(min_row=1,max_row=20,values_only=True),1):
                columns=[str(v).strip() if v is not None else '' for v in row]
                if all(column in columns for column in expected):
                    if any(columns.count(column)!=1 for column in expected):raise ValueError('Duplicate expected source header')
                    header=number;mapping=[columns.index(column) for column in expected];break
            if header is None:raise ValueError(f'Unrecognized header layout in {sheet.title}; prepare the documented schema first')
            matrix=[expected]
            for cells,formula_cells in zip(sheet.iter_rows(min_row=header+1),formula_sheet.iter_rows(min_row=header+1)):
                row=[]
                for index in mapping:
                    value=cells[index].value if index<len(cells) else None
                    formula=formula_cells[index] if index<len(formula_cells) else None
                    if value is None and formula is not None and formula.data_type=='f':raise ValueError(f'Formula has no cached result: {sheet.title}!{formula.coordinate}')
                    if isinstance(value,(datetime.datetime,datetime.date)):value={'date':value.isoformat()[:10]}
                    row.append(value)
                if any(value is not None for value in row):matrix.append(row)
            result[name]=matrix
        return result
    finally:values.close();formulas.close()

def main():
    parser=argparse.ArgumentParser(description=__doc__);parser.add_argument('input',type=Path);parser.add_argument('--output',type=Path,required=True);args=parser.parse_args()
    source=prepare(args.input)
    for name,columns in {'sw-details':[24,36,37,40,41],'hw-details':[17,29,55]}.items():
        for number,row in enumerate(source[name][1:],2):
            for column in columns:
                if row[column] is not None and not isinstance(row[column],(int,float)):raise ValueError(f'Expected numeric field in {name}, prepared row {number}, column {column+1}')
    args.output.parent.mkdir(parents=True,exist_ok=True);args.output.write_text(json.dumps(source,ensure_ascii=False,allow_nan=False,separators=(',',':')),encoding='utf8')
    print('Prepared '+str(args.output))

if __name__=='__main__':main()
