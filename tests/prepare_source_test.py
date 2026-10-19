import datetime, importlib.util, json, tempfile, unittest
from pathlib import Path
from openpyxl import Workbook

ROOT=Path(__file__).resolve().parents[1]
spec=importlib.util.spec_from_file_location('prepare_source',ROOT/'tools/prepare-source.py')
prepare_source=importlib.util.module_from_spec(spec);spec.loader.exec_module(prepare_source)

class PreparationTests(unittest.TestCase):
    def test_dates_and_empty_portfolio_preserve_schema(self):
        schema=json.loads((ROOT/'schemas/eix-columns.json').read_text())
        with tempfile.TemporaryDirectory() as folder:
            workbook=Workbook();workbook.remove(workbook.active)
            for name,columns in schema.items():
                sheet=workbook.create_sheet(name);sheet.append(columns)
                if name=='sw-details':
                    row=[None]*len(columns);row[4]='Example Entity';row[11]='Active';row[28]=datetime.date(2026,1,31);row[40]=0;sheet.append(row)
            file=Path(folder)/'input.xlsx';workbook.save(file)
            source=prepare_source.prepare(file)
            self.assertEqual(source['sw-details'][1][28],{'date':'2026-01-31'})
            self.assertEqual(source['sw-details'][1][40],0)
            self.assertEqual(len(source['hw-details']),1)

    def test_uncached_formula_fails_instead_of_disappearing(self):
        schema=json.loads((ROOT/'schemas/eix-columns.json').read_text())
        with tempfile.TemporaryDirectory() as folder:
            workbook=Workbook();workbook.remove(workbook.active)
            for name,columns in schema.items():
                sheet=workbook.create_sheet(name);sheet.append(columns)
                if name=='sw-details':
                    row=[None]*len(columns);row[40]='=1+1';sheet.append(row)
            file=Path(folder)/'input.xlsx';workbook.save(file)
            with self.assertRaisesRegex(ValueError,'cached result'):prepare_source.prepare(file)

if __name__=='__main__':unittest.main()
